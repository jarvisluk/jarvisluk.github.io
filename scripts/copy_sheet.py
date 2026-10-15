"""Copy review sheet <-> language packs.

    python3 scripts/copy_sheet.py export          # src/locales -> design/copy/文案对照.xlsx
    python3 scripts/copy_sheet.py import          # filled-in 修改后中文 / 修改后英文 -> src/locales
    python3 scripts/copy_sheet.py import --check  # list what would change, write nothing
    … --sheet other.xlsx                          # use another copy of the sheet

Needs openpyxl (pip install openpyxl). Export reads the packs through scripts/export-copy.mjs.
Import only applies rows whose current text still matches the packs, so an old sheet cannot
overwrite newer copy; rows from the 嵌入页面 sheet (public/embeds) are listed, never written.
"""
import glob, json, os, re, subprocess, sys
from html.parser import HTMLParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEET = os.path.join(ROOT, 'design/copy/文案对照.xlsx')
LOCALES = os.path.join(ROOT, 'src/locales')
LANGS = ('zh', 'en')
HEAD = ['#', '键名 Key / 位置', '中文（当前）', 'English (current)', '修改后中文', '修改后英文', '备注', '代码路径']
EMBEDS = '嵌入页面'
CJK = re.compile(r'[一-鿿]')


def embed_rows():
    """Text in public/embeds: data-zh on an element, English from data-en or the element's own text."""
    rows = []
    for f in sorted(glob.glob(f'{ROOT}/public/embeds/**/*.html', recursive=True)):
        rel = os.path.relpath(f, ROOT)

        class P(HTMLParser):
            def __init__(s):
                super().__init__(convert_charrefs=True); s.stack = []; s.found = []

            def handle_starttag(s, tag, attrs):
                a = dict(attrs)
                if tag in ('br', 'img', 'input', 'meta', 'link', 'hr', 'source', 'col', 'wbr'):
                    if 'data-zh' in a:
                        s.found.append((a['data-zh'], a.get('data-en') or a.get('aria-label') or a.get('placeholder') or '', s.getpos()[0]))
                    return
                s.stack.append([tag, a, [], s.getpos()[0]])

            def handle_endtag(s, tag):
                while s.stack:
                    t, a, txt, ln = s.stack.pop()
                    if s.stack: s.stack[-1][2].append(''.join(txt))
                    if 'data-zh' in a: s.found.append((a['data-zh'], a.get('data-en') or ' '.join(''.join(txt).split()), ln))
                    if t == tag: break

            def handle_data(s, d):
                if s.stack: s.stack[-1][2].append(d)

        src = open(f, encoding='utf-8').read()
        p = P(); p.feed(src)
        # Text set from scripts: T('中','en') / b(…) / bi(…), {zh:'…',en:'…'} and lang==='zh'?'…':'…'.
        q = r"'((?:[^'\\]|\\.)*)'"
        for ln, line in enumerate(src.splitlines(), 1):
            for pat in (r"\b(?:T|b|bi)\(" + q + r"\s*,\s*" + q + r"\)", r"zh:" + q + r"\s*,\s*en:" + q, r"lang===?'zh'\?" + q + ":" + q):
                for m in re.finditer(pat, line): p.found.append((m.group(1), m.group(2), ln))
        seen = set()
        for zh, en, ln in sorted(p.found, key=lambda x: x[2]):
            if (zh, en) in seen or not CJK.search(zh): continue
            seen.add((zh, en)); rows.append((f'{rel}:{ln}', zh, en, rel))
    return rows


def export():
    rows = json.loads(subprocess.run(['node', 'scripts/export-copy.mjs'], cwd=ROOT, check=True, capture_output=True, text=True).stdout)
    sheets = {}
    for r in rows: sheets.setdefault(r['sheet'], []).append((r['key'], r['zh'], r['en'], 'src/locales/*.json'))
    sheets[EMBEDS] = embed_rows()
    wb = Workbook(); info = wb.active; info.title = '说明'
    counts = []
    for name, items in sheets.items():
        ws = wb.create_sheet(name[:31]); ws.append(HEAD)
        for n, (key, zh, en, where) in enumerate(items, 1): ws.append([n, key, zh, en, None, None, None, where])
        for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='EEEEEE')
        for col, w in zip('ABCDEFGH', [5, 34, 50, 50, 36, 36, 20, 30]): ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'C2'; counts.append((name, len(items)))
    notes = [
        '网站全部文案（中英对照），从语言包 src/locales/zh.json、en.json 导出，按页面阅读顺序排列。',
        '在「修改后中文」「修改后英文」两列填写新文案；不需要改的留空即可。填好后运行 npm run copy:import 写回语言包。',
        '「备注」可写想法，比如「删掉」「合并到上一条」「语气更口语」；备注不会自动生效，导入时会列出来。',
        '「键名」就是语言包里的键，由页面位置组成：项目 id、章节 id、模块序号、图片 id 等。请不要改键名和「中文（当前）」「English (current)」两列，导入靠它们核对。',
        '带 {n}、{date} 这类花括号的是占位符，会替换成数字或日期，修改时请保留。',
        '「嵌入页面」是 public/embeds 下独立的可交互演示页，文字写在各自的 HTML 里，不在语言包中；这一页的修改导入时只会列出，需要手动改。',
        '', '各表条数：'] + [f'{n}：{c}' for n, c in counts] + [f'合计：{sum(c for _, c in counts)}']
    for t in notes: info.append([t])
    info.column_dimensions['A'].width = 110
    wb.save(SHEET)
    print(f'Exported {sum(c for _, c in counts)} rows to {os.path.relpath(SHEET, ROOT)}')


def lookup(pack, key):
    node = pack
    for k in key.split('.'):
        if isinstance(node, list):
            if not k.isdigit() or int(k) >= len(node): return None
            node = node[int(k)]
        elif isinstance(node, dict) and k in node:
            node = node[k]
        else:
            return None
    return node


def assign(pack, key, value):
    *path, last = key.split('.')
    node = pack
    for k in path: node = node[int(k)] if isinstance(node, list) else node[k]
    if isinstance(node, list): node[int(last)] = value
    else: node[last] = value


def text(v):
    return '' if v is None else str(v)


def import_(check):
    packs = {l: json.load(open(f'{LOCALES}/{l}.json', encoding='utf-8')) for l in LANGS}
    wb = load_workbook(SHEET, read_only=True)
    applied, problems, embeds, notes = [], [], [], []
    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(min_row=2, values_only=True):
            row = (list(row) + [None] * 8)[:8]
            _, key, cur_zh, cur_en, new_zh, new_en, note, _ = row
            key = text(key).strip()
            new = {'zh': text(new_zh).strip(), 'en': text(new_en).strip()}
            if note: notes.append(f'{ws.title} · {key}：{note}')
            if not (new['zh'] or new['en']): continue
            if ws.title == EMBEDS:
                embeds.append(f'{key}  {new["zh"] or "-"} / {new["en"] or "-"}'); continue
            current = {'zh': text(cur_zh), 'en': text(cur_en)}
            for lang in LANGS:
                if not new[lang]: continue
                live = lookup(packs[lang], key)
                if not isinstance(live, str):
                    problems.append(f'{key} [{lang}]: key not in the language pack'); continue
                if live != current[lang]:
                    problems.append(f'{key} [{lang}]: pack changed since export, skipped (re-export first)'); continue
                if new[lang] == live: continue
                holes = sorted(re.findall(r'\{\w+\}', live))
                if holes != sorted(re.findall(r'\{\w+\}', new[lang])):
                    problems.append(f'{key} [{lang}]: keep the placeholders {" ".join(holes)}, skipped'); continue
                assign(packs[lang], key, new[lang]); applied.append(f'{key} [{lang}]: {live} → {new[lang]}')
    print(f'{"Would apply" if check else "Applied"} {len(applied)} change(s):'); [print('  ' + a) for a in applied]
    if problems: print(f'\nSkipped {len(problems)}:'); [print('  ' + p) for p in problems]
    if embeds: print(f'\nEmbedded pages, edit by hand ({len(embeds)}):'); [print('  ' + e) for e in embeds]
    if notes: print(f'\nNotes ({len(notes)}):'); [print('  ' + n) for n in notes]
    if applied and not check:
        for l in LANGS:
            with open(f'{LOCALES}/{l}.json', 'w', encoding='utf-8') as f: f.write(json.dumps(packs[l], ensure_ascii=False, indent=1) + '\n')
        print('\nRun `python3 scripts/copy_sheet.py export` to refresh the sheet once you are done.')
    return 1 if problems else 0


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else ''
    if '--sheet' in sys.argv: SHEET = os.path.abspath(sys.argv[sys.argv.index('--sheet') + 1])
    if cmd == 'export': export()
    elif cmd == 'import': sys.exit(import_('--check' in sys.argv))
    else: sys.exit(__doc__)
