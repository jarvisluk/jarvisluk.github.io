#!/usr/bin/env python3
"""Build a page-oriented source library from the PolyU handoff folder."""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import fitz
from PIL import Image
from pptx import Presentation


SOURCE = Path("/Users/junrong/Desktop/PolyU Final")
ROOT = Path("content-library/polyu-final")

PROJECTS = {
    "Bambu Lab Project.pdf": ("bambu-lab-flagship-store", "Bambu Lab Flagship Store"),
    "SD5431_Group 2_ Final Presentation.pdf": ("sf-express-eating-habits", "SF Express Eating Habits"),
    "SD5413 Mirs Bay Final Presentation_GeoGroup.pdf": ("mirs-bay-geopark", "Mirs Bay x Geopark"),
    "VA final.pdf": ("va-catering-feedback", "VA Catering Feedback System"),
    "SD5432 - FINAL PPP.pptx": ("guangzhou-ebike-parking", "Guangzhou E-bike Parking & Pick-up"),
}

VIDEOS = {
    "Bambu Demo.mov": "bambu-lab-flagship-store",
    "Bambu Workshop.mp4": "bambu-lab-flagship-store",
}

EXTRA_PROJECTS = {
    Path("/Users/junrong/Desktop/毕设开题汇报.pdf"): ("pawpals-tnr", "PawPals - TNR Stray Cat Management System"),
    Path("/Users/junrong/Desktop/Just A 2nd.pdf"): ("just-a-2nd", "Just A 2nd"),
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean_text(text: str) -> str:
    return "\n".join(line.rstrip() for line in text.replace("\u2028", "\n").splitlines()).strip() + "\n"


def title_from(text: str, fallback: str) -> str:
    for line in text.splitlines():
        line = " ".join(line.split())
        if line:
            return line[:160]
    return fallback


def ffprobe(path: Path) -> dict:
    cmd = [
        "ffprobe", "-v", "error", "-show_entries",
        "format=duration,size,format_name:stream=index,codec_type,codec_name,width,height,r_frame_rate",
        "-of", "json", str(path),
    ]
    return json.loads(subprocess.check_output(cmd, text=True))


def extract_pdf(path: Path, slug: str, project_title: str) -> dict:
    project = ROOT / "projects" / slug
    pages_root = project / "pages"
    pages_root.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, project / "source" / path.name) if (project / "source").exists() else None
    (project / "source").mkdir(exist_ok=True)
    shutil.copy2(path, project / "source" / path.name)
    doc = fitz.open(path)
    records = []
    for index, page in enumerate(doc, 1):
        page_dir = pages_root / f"page-{index:03d}"
        images_dir = page_dir / "images"
        images_dir.mkdir(parents=True, exist_ok=True)
        text = clean_text(page.get_text("text"))
        (page_dir / "text.md").write_text(text, encoding="utf-8")

        scale = 1600 / page.rect.width
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        preview = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        preview.save(page_dir / "page-preview.webp", "WEBP", quality=88, method=6)

        image_records = []
        seen = set()
        for occurrence, info in enumerate(page.get_image_info(xrefs=True), 1):
            xref = info.get("xref", 0)
            if not xref:
                continue
            key = (xref, tuple(round(v, 3) for v in info.get("bbox", ())))
            if key in seen:
                continue
            seen.add(key)
            raw = doc.extract_image(xref)
            ext = raw.get("ext", "bin")
            name = f"image-{occurrence:03d}-xref-{xref}.{ext}"
            (images_dir / name).write_bytes(raw["image"])
            image_records.append({
                "file": f"images/{name}",
                "xref": xref,
                "bbox": [round(v, 3) for v in info.get("bbox", ())],
                "pixel_width": info.get("width"),
                "pixel_height": info.get("height"),
                "has_mask": bool(raw.get("smask")),
            })
        record = {
            "page": index,
            "title": title_from(text, f"Page {index}"),
            "text_file": "text.md",
            "preview_file": "page-preview.webp",
            "images": image_records,
        }
        (page_dir / "page.json").write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
        records.append(record)

    manifest = {
        "slug": slug,
        "title": project_title,
        "source_file": path.name,
        "source_type": "pdf",
        "page_count": len(doc),
        "pages": records,
    }
    (project / "project.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def pptx_media_for_slide(presentation: Path, slide_num: int) -> list[str]:
    slide = f"ppt/slides/slide{slide_num}.xml"
    rels = f"ppt/slides/_rels/slide{slide_num}.xml.rels"
    with zipfile.ZipFile(presentation) as zf:
        if rels not in zf.namelist():
            return []
        import xml.etree.ElementTree as ET
        root = ET.fromstring(zf.read(rels))
        targets = []
        for rel in root:
            target = rel.attrib.get("Target", "")
            if "../media/" in target:
                targets.append("ppt/media/" + target.split("../media/", 1)[1])
        return list(dict.fromkeys(targets))


def extract_pptx(path: Path, slug: str, project_title: str) -> dict:
    project = ROOT / "projects" / slug
    pages_root = project / "pages"
    source_dir = project / "source"
    pages_root.mkdir(parents=True, exist_ok=True)
    source_dir.mkdir(exist_ok=True)
    shutil.copy2(path, source_dir / path.name)
    prs = Presentation(path)
    records = []
    with zipfile.ZipFile(path) as zf:
        for index, slide in enumerate(prs.slides, 1):
            page_dir = pages_root / f"page-{index:03d}"
            images_dir = page_dir / "images"
            images_dir.mkdir(parents=True, exist_ok=True)
            chunks = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    chunks.append(shape.text.strip())
            text = clean_text("\n\n".join(chunks))
            (page_dir / "text.md").write_text(text, encoding="utf-8")
            image_records = []
            for occurrence, member in enumerate(pptx_media_for_slide(path, index), 1):
                ext = Path(member).suffix.lower() or ".bin"
                name = f"image-{occurrence:03d}{ext}"
                data = zf.read(member)
                (images_dir / name).write_bytes(data)
                image_records.append({"file": f"images/{name}", "package_member": member, "bytes": len(data)})
            record = {
                "page": index,
                "title": title_from(text, f"Slide {index}"),
                "text_file": "text.md",
                "preview_file": None,
                "images": image_records,
            }
            (page_dir / "page.json").write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
            records.append(record)
    manifest = {
        "slug": slug,
        "title": project_title,
        "source_file": path.name,
        "source_type": "pptx",
        "page_count": len(prs.slides),
        "pages": records,
        "note": "Slide text and directly related package media were extracted. Preview files are added when LibreOffice rendering is available.",
    }
    (project / "project.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def add_pptx_previews(path: Path, slug: str) -> bool:
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        return False
    temp = ROOT / ".render"
    temp.mkdir(exist_ok=True)
    result = subprocess.run(
        [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    if result.returncode:
        shutil.rmtree(temp)
        return False
    pdf = temp / (path.stem + ".pdf")
    doc = fitz.open(pdf)
    project = ROOT / "projects" / slug
    manifest_path = project / "project.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for index, page in enumerate(doc, 1):
        page_dir = project / "pages" / f"page-{index:03d}"
        scale = 1600 / page.rect.width
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(page_dir / "page-preview.webp", "WEBP", quality=88, method=6)
        rec_path = page_dir / "page.json"
        record = json.loads(rec_path.read_text(encoding="utf-8"))
        record["preview_file"] = "page-preview.webp"
        rec_path.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
        manifest["pages"][index - 1]["preview_file"] = "page-preview.webp"
    manifest["note"] = "Slide text, page previews, and directly related package media were extracted."
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    shutil.rmtree(temp)
    return True


def main() -> None:
    if ROOT.exists():
        shutil.rmtree(ROOT)
    (ROOT / "projects").mkdir(parents=True)
    manifests = []
    for filename, (slug, title) in PROJECTS.items():
        path = SOURCE / filename
        manifests.append(extract_pdf(path, slug, title) if path.suffix.lower() == ".pdf" else extract_pptx(path, slug, title))
        if path.suffix.lower() == ".pptx":
            add_pptx_previews(path, slug)
    for path, (slug, title) in EXTRA_PROJECTS.items():
        manifests.append(extract_pdf(path, slug, title))

    video_records = []
    for filename, slug in VIDEOS.items():
        source = SOURCE / filename
        target_dir = ROOT / "projects" / slug / "videos"
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / filename
        shutil.copy2(source, target)
        metadata = ffprobe(target)
        poster = target.with_suffix(".poster.jpg")
        subprocess.run(["ffmpeg", "-y", "-ss", "00:00:02", "-i", str(target), "-frames:v", "1", "-vf", "scale=1280:-2", str(poster)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)
        rec = {"project": slug, "file": f"projects/{slug}/videos/{filename}", "poster": f"projects/{slug}/videos/{poster.name}", "sha256": sha256(target), "metadata": metadata}
        video_records.append(rec)
    bambu_manifest_path = ROOT / "projects" / "bambu-lab-flagship-store" / "project.json"
    bambu = json.loads(bambu_manifest_path.read_text(encoding="utf-8"))
    bambu["videos"] = [v for v in video_records if v["project"] == "bambu-lab-flagship-store"]
    bambu_manifest_path.write_text(json.dumps(bambu, ensure_ascii=False, indent=2), encoding="utf-8")

    index = {"source": str(SOURCE), "projects": [{"slug": m["slug"], "title": m["title"], "source_file": m["source_file"], "page_count": m["page_count"]} for m in manifests], "videos": video_records}
    (ROOT / "library.json").write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    readme = "# PolyU Final 素材库\n\n每个项目位于 `projects/<项目>/`。每页位于 `pages/page-NNN/`，包含：\n\n- `text.md`：该页提取的文字\n- `page-preview.webp`：该页完整预览（若渲染器可用）\n- `images/`：该页嵌入的原始图片\n- `page.json`：页标题、图片位置及尺寸等机器可读信息\n\n项目根目录的 `project.json` 是项目清单；`library.json` 是总清单。视频存放于对应项目的 `videos/`，并带有 poster 图。\n"
    (ROOT / "README.md").write_text(readme, encoding="utf-8")


if __name__ == "__main__":
    main()
